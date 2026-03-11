import os
import re
import time
import json
import logging
from io import BytesIO
from contextvars import ContextVar
from dataclasses import dataclass
from typing import Optional, Any
from uuid import uuid4
from fastapi import FastAPI, Request, UploadFile, File, Form
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from dotenv import load_dotenv

from agent_framework import ChatAgent, ChatMessage, ConcurrentBuilder, AgentRunUpdateEvent, WorkflowOutputEvent, GroupChatBuilder
from agent_framework.azure import AzureOpenAIChatClient
from agent_framework.openai import OpenAIChatClient
# from agent_framework.azure import AzureAISearchContextProvider
from typing import Annotated
from pydantic import Field
from azure.search.documents import SearchClient
from azure.core.credentials import AzureKeyCredential
from azure.search.documents.models import QueryType
from session_state import SessionStore, UploadedDocument
from translations import get_text

load_dotenv()

# Logging configuration
logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger(__name__)

app = FastAPI()

# CORS configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5000",   # Flask / .NET
        "http://localhost:5001",   # Flask
        "https://localhost:5001",  # .NET (HTTPS)
        "http://localhost:8501",   # Streamlit
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Language setting
LANGUAGE = os.getenv("LANGUAGE", "ja")  # Default: Japanese

# Azure AI Search settings
SEARCH_ENDPOINT = os.getenv("SEARCH_ENDPOINT")
SEARCH_API_KEY = os.getenv("SEARCH_API_KEY")
SEARCH_INDEX_NAME = os.getenv("SEARCH_INDEX_NAME")
SEARCH_SEMANTIC_CONFIG = os.getenv("SEARCH_SEMANTIC_CONFIG", "default")
SESSION_HISTORY_MESSAGES = max(int(os.getenv("SESSION_HISTORY_MESSAGES", "20")), 4)
PROMPT_HISTORY_MESSAGES = max(int(os.getenv("PROMPT_HISTORY_MESSAGES", "8")), 2)
PLAN_STEP_LIMIT = max(int(os.getenv("PLAN_STEP_LIMIT", "4")), 2)
PLAN_TOOL_LIMIT = max(int(os.getenv("PLAN_TOOL_LIMIT", "4")), 1)
MODE_GENERAL = "general"
MODE_GUIDELINE = "guideline"
MODE_IDOBATA = "idobata"

session_store = SessionStore(
    max_messages=SESSION_HISTORY_MESSAGES,
    store_dir=os.getenv("BACKEND_SESSION_STORE_PATH", "data/backend_sessions"),
)
SEARCH_EXCERPT_LIMIT = max(int(os.getenv("SEARCH_EXCERPT_LIMIT", "280")), 80)
SEARCH_TRACE_CONTEXT: ContextVar[list[dict[str, Any]] | None] = ContextVar("search_trace_context", default=None)
MAX_UPLOAD_FILES = max(int(os.getenv("MAX_UPLOAD_FILES", "5")), 1)
MAX_UPLOAD_FILE_SIZE_BYTES = max(int(os.getenv("MAX_UPLOAD_FILE_SIZE_BYTES", "5242880")), 1024)
MAX_UPLOAD_TEXT_CHARS = max(int(os.getenv("MAX_UPLOAD_TEXT_CHARS", "12000")), 1000)
MAX_PROMPT_DOCUMENT_CHARS = max(int(os.getenv("MAX_PROMPT_DOCUMENT_CHARS", "8000")), 1000)

ALLOWED_UPLOAD_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".pdf",
    ".docx",
}

MODEL_LABEL_DEFAULTS = {
    "gpt-4.1": "GPT-4.1",
    "gpt-4.1-mini": "GPT-4.1 Mini",
    "gpt-4.1-nano": "GPT-4.1 Nano",
    "gpt-4-1": "GPT-4.1",
    "gpt-4-1-mini": "GPT-4.1 Mini",
    "gpt-4-1-nano": "GPT-4.1 Nano",
    "gpt-5": "GPT-5",
    "gpt-5-mini": "GPT-5 Mini",
    "gpt-5-nano": "GPT-5 Nano",
    "gpt-5.2": "GPT-5.2",
    "gpt-5.2-chat": "GPT-5.2 Chat",
}
MODEL_PROVIDER_AZURE = "azure"
MODEL_PROVIDER_OPENAI = "openai"
MODEL_PROVIDER_LABELS = {
    MODEL_PROVIDER_AZURE: "Azure OpenAI",
    MODEL_PROVIDER_OPENAI: "OpenAI",
}


@dataclass(frozen=True)
class ModelConfig:
    model_id: str
    label: str
    provider: str
    api_key: str
    endpoint: str | None = None
    deployment_name: str | None = None
    api_version: str | None = None
    provider_model_id: str | None = None
    org_id: str | None = None
    approval_required: bool = False
    approval_reason: str | None = None

    @property
    def provider_label(self) -> str:
        return MODEL_PROVIDER_LABELS.get(self.provider, self.provider)

    @property
    def target_name(self) -> str:
        return self.deployment_name or self.provider_model_id or self.model_id


def _parse_csv_env(name: str) -> list[str]:
    raw_value = os.getenv(name, "")
    return [item.strip() for item in raw_value.split(",") if item.strip()]


def _parse_bool_env(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def _model_env_suffix(model_name: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", model_name).strip("_").upper()
    return normalized or "DEFAULT"


def _model_label_candidates(model_name: str | None) -> list[str]:
    if not model_name:
        return []

    candidates: list[str] = []

    def add_candidate(value: str) -> None:
        if value and value not in candidates:
            candidates.append(value)

    add_candidate(model_name)
    add_candidate(model_name.replace("-", "."))
    add_candidate(model_name.replace(".", "-"))

    provider_prefixes = (f"{MODEL_PROVIDER_AZURE}-", f"{MODEL_PROVIDER_OPENAI}-")
    for prefix in provider_prefixes:
        if model_name.startswith(prefix):
            bare_name = model_name[len(prefix):]
            add_candidate(bare_name)
            add_candidate(bare_name.replace("-", "."))
            add_candidate(bare_name.replace(".", "-"))
            break

    return candidates


def _resolve_model_label(model_name: str | None) -> str | None:
    for candidate in _model_label_candidates(model_name):
        label = MODEL_LABEL_DEFAULTS.get(candidate)
        if label:
            return label
    return None


def _default_model_label(model_id: str, provider: str, reference_name: str | None = None) -> str:
    base_name = _resolve_model_label(model_id)
    if not base_name and reference_name:
        base_name = _resolve_model_label(reference_name)
    if not base_name:
        base_name = model_id
    return f"{base_name} ({MODEL_PROVIDER_LABELS.get(provider, provider)})"


def _warn_incomplete_model_config(model_id: str, prefix: str, configured_values: dict[str, str | None], required_keys: list[str]) -> None:
    if not any(value for value in configured_values.values()):
        return

    missing = [key for key in required_keys if not configured_values.get(key)]
    if missing:
        logger.warning(
            "Skipping model '%s': missing %s in env prefix %s",
            model_id,
            ", ".join(missing),
            prefix,
        )


def _build_model_config(model_id: str) -> ModelConfig | None:
    suffix = _model_env_suffix(model_id)
    prefix = f"LLM_MODEL_{suffix}_"

    provider = os.getenv(f"{prefix}PROVIDER", "").strip().lower()
    api_key = os.getenv(f"{prefix}API_KEY", "").strip()
    endpoint = os.getenv(f"{prefix}ENDPOINT", "").strip()
    base_url = os.getenv(f"{prefix}BASE_URL", "").strip()
    api_version = os.getenv(f"{prefix}API_VERSION", "").strip() or None
    deployment_name = os.getenv(f"{prefix}DEPLOYMENT", "").strip() or None
    provider_model_id = os.getenv(f"{prefix}MODEL_ID", "").strip() or None
    org_id = os.getenv(f"{prefix}ORG_ID", "").strip() or None
    label = os.getenv(f"{prefix}LABEL", "").strip()
    approval_required = _parse_bool_env(f"{prefix}REQUIRES_APPROVAL", False)
    approval_reason = os.getenv(f"{prefix}APPROVAL_REASON", "").strip() or None

    if provider == MODEL_PROVIDER_AZURE:
        configured_values = {
            "PROVIDER": provider,
            "API_KEY": api_key,
            "ENDPOINT": endpoint,
            "DEPLOYMENT": deployment_name,
            "API_VERSION": api_version,
        }
        if api_key and endpoint and deployment_name:
            return ModelConfig(
                model_id=model_id,
                label=label or _default_model_label(model_id, provider),
                provider=provider,
                api_key=api_key,
                endpoint=endpoint,
                deployment_name=deployment_name,
                api_version=api_version,
                approval_required=approval_required,
                approval_reason=approval_reason,
            )
        _warn_incomplete_model_config(model_id, prefix, configured_values, ["API_KEY", "ENDPOINT", "DEPLOYMENT"])
        return None

    if provider == MODEL_PROVIDER_OPENAI:
        resolved_endpoint = endpoint or base_url or None
        configured_values = {
            "PROVIDER": provider,
            "API_KEY": api_key,
            "MODEL_ID": provider_model_id,
            "ENDPOINT": resolved_endpoint,
            "ORG_ID": org_id,
        }
        if api_key and provider_model_id:
            return ModelConfig(
                model_id=model_id,
                label=label or _default_model_label(model_id, provider, reference_name=provider_model_id),
                provider=provider,
                api_key=api_key,
                endpoint=resolved_endpoint,
                provider_model_id=provider_model_id,
                org_id=org_id,
                approval_required=approval_required,
                approval_reason=approval_reason,
            )
        _warn_incomplete_model_config(model_id, prefix, configured_values, ["API_KEY", "MODEL_ID"])
        return None

    configured_values = {
        "PROVIDER": provider,
        "API_KEY": api_key,
        "ENDPOINT": endpoint or base_url,
        "API_VERSION": api_version,
        "DEPLOYMENT": deployment_name,
        "MODEL_ID": provider_model_id,
        "ORG_ID": org_id,
        "LABEL": label,
    }
    if any(value for value in configured_values.values()):
        logger.warning(
            "Skipping model '%s': unsupported or missing PROVIDER in env prefix %s",
            model_id,
            prefix,
        )
    return None


def _build_legacy_azure_model_config(model_id: str) -> ModelConfig | None:
    suffix = _model_env_suffix(model_id)
    prefix = f"AZURE_OPENAI_MODEL_{suffix}_"

    api_key = os.getenv(f"{prefix}API_KEY", "").strip()
    endpoint = os.getenv(f"{prefix}ENDPOINT", "").strip()
    deployment_name = os.getenv(f"{prefix}DEPLOYMENT", "").strip()
    api_version = os.getenv(f"{prefix}API_VERSION", "").strip() or None
    label = os.getenv(f"{prefix}LABEL", "").strip()

    configured_values = {
        "API_KEY": api_key,
        "ENDPOINT": endpoint,
        "DEPLOYMENT": deployment_name,
        "API_VERSION": api_version,
    }
    if api_key and endpoint and deployment_name:
        return ModelConfig(
            model_id=model_id,
            label=label or _default_model_label(model_id, MODEL_PROVIDER_AZURE),
            provider=MODEL_PROVIDER_AZURE,
            api_key=api_key,
            endpoint=endpoint,
            deployment_name=deployment_name,
            api_version=api_version,
        )
    _warn_incomplete_model_config(model_id, prefix, configured_values, ["API_KEY", "ENDPOINT", "DEPLOYMENT"])
    return None


def _load_legacy_model_configs() -> list[ModelConfig]:
    api_key = os.getenv("AZURE_OPENAI_API_KEY", "").strip()
    endpoint = os.getenv("AZURE_OPENAI_ENDPOINT", "").strip()
    api_version = os.getenv("AZURE_OPENAI_API_VERSION", "").strip() or None
    default_deployment = os.getenv("AZURE_OPENAI_DEPLOYMENT", "").strip()

    if not api_key or not endpoint:
        return []

    deployment_map = {
        "gpt-4.1": os.getenv("AZURE_OPENAI_DEPLOYMENT_GPT41", "").strip() or default_deployment,
        "gpt-4.1-mini": os.getenv("AZURE_OPENAI_DEPLOYMENT_GPT41_MINI", "").strip() or default_deployment,
        "gpt-4.1-nano": os.getenv("AZURE_OPENAI_DEPLOYMENT_GPT41_NANO", "").strip() or default_deployment,
        "gpt-5.2": os.getenv("AZURE_OPENAI_DEPLOYMENT_GPT52", "").strip() or default_deployment,
        "gpt-5.2-chat": os.getenv("AZURE_OPENAI_DEPLOYMENT_GPT52_CHAT", "").strip() or default_deployment,
    }

    configs = []
    for model_name, deployment_name in deployment_map.items():
        if not deployment_name:
            continue
        configs.append(
            ModelConfig(
                model_id=model_name,
                label=_default_model_label(model_name, MODEL_PROVIDER_AZURE),
                provider=MODEL_PROVIDER_AZURE,
                api_key=api_key,
                endpoint=endpoint,
                deployment_name=deployment_name,
                api_version=api_version,
            )
        )
    return configs


def _load_model_registry() -> tuple[dict[str, ModelConfig], str | None]:
    configured_model_names = _parse_csv_env("LLM_MODELS")
    legacy_named_models = _parse_csv_env("AZURE_OPENAI_MODELS")

    registry: dict[str, ModelConfig] = {}
    if configured_model_names:
        for model_id in configured_model_names:
            config = _build_model_config(model_id)
            if config is not None:
                registry[model_id] = config
    elif legacy_named_models:
        for model_id in legacy_named_models:
            config = _build_legacy_azure_model_config(model_id)
            if config is not None:
                registry[model_id] = config
    else:
        for config in _load_legacy_model_configs():
            registry[config.model_id] = config

    default_model_name = os.getenv("DEFAULT_MODEL", "").strip()
    if default_model_name and default_model_name not in registry:
        logger.warning("DEFAULT_MODEL '%s' is not configured. Falling back to the first available model.", default_model_name)
        default_model_name = ""

    if not default_model_name and registry:
        default_model_name = next(iter(registry))

    return registry, default_model_name or None


MODEL_REGISTRY, DEFAULT_MODEL_NAME = _load_model_registry()
if not MODEL_REGISTRY:
    logger.warning("No models are configured. Set LLM_MODELS and model-specific env vars, or use the legacy Azure fallback.")


def _build_chat_client(config: ModelConfig):
    if config.provider == MODEL_PROVIDER_AZURE:
        return AzureOpenAIChatClient(
            api_key=config.api_key,
            endpoint=config.endpoint,
            deployment_name=config.deployment_name,
            api_version=config.api_version,
        )

    if config.provider == MODEL_PROVIDER_OPENAI:
        kwargs = {
            "api_key": config.api_key,
            "model_id": config.provider_model_id,
        }
        if config.endpoint:
            kwargs["base_url"] = config.endpoint
        if config.org_id:
            kwargs["org_id"] = config.org_id
        return OpenAIChatClient(**kwargs)

    raise ValueError(f"Unsupported provider: {config.provider}")


MODEL_CLIENTS = {
    model_id: _build_chat_client(config)
    for model_id, config in MODEL_REGISTRY.items()
}

def get_chat_client_for_model(model_name: str = None):
    """Return a chat client for the requested model."""
    if not MODEL_CLIENTS:
        return None

    resolved_model_name = model_name if model_name in MODEL_CLIENTS else DEFAULT_MODEL_NAME
    if not resolved_model_name:
        return None

    config = MODEL_REGISTRY[resolved_model_name]
    logger.info(
        get_text(
            'log_model_selected',
            LANGUAGE,
            model=resolved_model_name,
            provider=config.provider_label,
            target=config.target_name,
        )
    )
    return MODEL_CLIENTS[resolved_model_name]


def get_requested_model_name(body: dict) -> str | None:
    raw_model_name = body.get("model", "")
    if raw_model_name is None:
        return DEFAULT_MODEL_NAME
    model_name = str(raw_model_name).strip()
    if model_name:
        return model_name
    return DEFAULT_MODEL_NAME


def get_model_metadata() -> list[dict[str, object]]:
    return [
        {
            "id": config.model_id,
            "label": config.label,
            "provider": config.provider,
            "provider_label": config.provider_label,
            "target": config.target_name,
            "is_default": config.model_id == DEFAULT_MODEL_NAME,
            "approval_required": config.approval_required,
            "approval_reason": config.approval_reason,
        }
        for config in MODEL_REGISTRY.values()
    ]


def get_session_id(body: dict, default: str = "default") -> str:
    raw_session_id = body.get("session_id", default)
    if raw_session_id is None:
        return default
    session_id = str(raw_session_id).strip()
    return session_id or default


def _message_role_value(message: ChatMessage) -> str:
    role = getattr(message, "role", "")
    return getattr(role, "value", str(role))


def _format_history_messages(messages: list[ChatMessage]) -> str:
    formatted = []
    for message in messages:
        text = (message.text or "").strip()
        if not text:
            continue
        role = _message_role_value(message)
        if role == "system":
            continue
        speaker = message.author_name or ("User" if role == "user" else "Assistant")
        collapsed_text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
        formatted.append(f"{speaker}: {collapsed_text[:1200]}")
    return "\n\n".join(formatted)


def _coerce_plan_text(value: object, limit: int = 240) -> str:
    if value is None:
        return ""

    if isinstance(value, dict):
        parts = []
        for key in ("title", "step", "name", "description", "detail"):
            item = value.get(key)
            if item:
                parts.append(str(item).strip())
        text = " - ".join(parts) if parts else json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)

    text = re.sub(r"\s+", " ", text).strip(" -\n\t")
    return text[:limit]


def _compact_text(value: object, limit: int = 240) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit]


def _append_search_trace(trace: dict[str, Any]) -> None:
    collector = SEARCH_TRACE_CONTEXT.get()
    if collector is not None:
        collector.append(trace)


def _collect_evidence_from_traces(traces: list[dict[str, Any]]) -> list[dict[str, str]]:
    evidence_items: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    for trace in traces:
        for item in trace.get("results", []):
            source = _compact_text(item.get("source"), limit=160)
            excerpt = _compact_text(item.get("excerpt"), limit=SEARCH_EXCERPT_LIMIT)
            if not source or not excerpt:
                continue
            key = (source, excerpt)
            if key in seen:
                continue
            seen.add(key)
            evidence_items.append({"source": source, "excerpt": excerpt})

    return evidence_items


def _guess_upload_extension(filename: str) -> str:
    _, extension = os.path.splitext(filename or "")
    return extension.lower()


def _decode_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp932", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_document_text(filename: str, data: bytes) -> str:
    extension = _guess_upload_extension(filename)

    if extension in {".txt", ".md", ".csv", ".json"}:
        return _decode_text_bytes(data)

    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        pages = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(page_text)
        return "\n\n".join(pages)

    if extension == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n".join(paragraphs)

    raise ValueError(get_text('upload_error_unsupported_type', LANGUAGE, filename=filename))


def _document_metadata(document: UploadedDocument) -> dict[str, Any]:
    return {
        "id": document.file_id,
        "name": document.name,
        "content_type": document.content_type,
        "size_bytes": document.size_bytes,
        "preview_text": document.preview_text,
        "text_length": len(document.extracted_text),
    }


def _format_uploaded_documents(documents: list[UploadedDocument]) -> str:
    if not documents:
        return ""

    parts: list[str] = []
    remaining_chars = MAX_PROMPT_DOCUMENT_CHARS
    for index, document in enumerate(documents, start=1):
        if remaining_chars <= 0:
            break
        available_chars = max(200, remaining_chars)
        excerpt = document.extracted_text[:available_chars].strip()
        if not excerpt:
            continue
        if len(document.extracted_text) > len(excerpt):
            excerpt = f"{excerpt}\n...[truncated]"
        block = f"[Document {index}] {document.name}\n{excerpt}"
        parts.append(block)
        remaining_chars -= len(excerpt)

    if not parts:
        return ""

    return "Uploaded session documents:\n\n" + "\n\n".join(parts)


def _normalize_plan_items(values: object, *, max_items: int, fallback_prefix: str) -> list[str]:
    candidates: list[object]
    if isinstance(values, list):
        candidates = values
    elif values:
        candidates = [values]
    else:
        candidates = []

    normalized: list[str] = []
    for item in candidates:
        text = _coerce_plan_text(item)
        if not text or text in normalized:
            continue
        normalized.append(text)
        if len(normalized) >= max_items:
            break

    if normalized:
        return normalized

    fallback_count = max(1, min(max_items, 2))
    return [f"{fallback_prefix} {index}" for index in range(1, fallback_count + 1)]


def _fallback_execution_plan(prompt: str) -> dict[str, object]:
    request_summary = _coerce_plan_text(prompt, limit=160) or get_text('plan_fallback_goal', LANGUAGE)
    return {
        "goal": request_summary,
        "steps": [
            get_text('plan_fallback_step_understand', LANGUAGE),
            get_text('plan_fallback_step_reason', LANGUAGE),
            get_text('plan_fallback_step_answer', LANGUAGE),
        ],
        "tools": [get_text('plan_fallback_tool_model', LANGUAGE)],
        "completion_criteria": [
            get_text('plan_fallback_completion_relevance', LANGUAGE),
            get_text('plan_fallback_completion_actionable', LANGUAGE),
        ],
    }


def _normalize_execution_plan(raw_plan: object, prompt: str) -> dict[str, object]:
    if not isinstance(raw_plan, dict):
        return _fallback_execution_plan(prompt)

    goal = _coerce_plan_text(raw_plan.get("goal"), limit=200) or _coerce_plan_text(prompt, limit=160)
    steps = _normalize_plan_items(
        raw_plan.get("steps"),
        max_items=PLAN_STEP_LIMIT,
        fallback_prefix=get_text('plan_fallback_step_label', LANGUAGE),
    )
    tools = _normalize_plan_items(
        raw_plan.get("tools") or [get_text('plan_fallback_tool_model', LANGUAGE)],
        max_items=PLAN_TOOL_LIMIT,
        fallback_prefix=get_text('plan_fallback_tool_label', LANGUAGE),
    )
    completion_criteria = _normalize_plan_items(
        raw_plan.get("completion_criteria"),
        max_items=3,
        fallback_prefix=get_text('plan_fallback_completion_label', LANGUAGE),
    )

    return {
        "goal": goal,
        "steps": steps,
        "tools": tools,
        "completion_criteria": completion_criteria,
    }


def _extract_plan_payload(planner_output: str, prompt: str) -> dict[str, object]:
    planner_output = planner_output.strip()
    if not planner_output:
        return _fallback_execution_plan(prompt)

    decoder = json.JSONDecoder()
    for index, char in enumerate(planner_output):
        if char != "{":
            continue
        try:
            raw_plan, _ = decoder.raw_decode(planner_output[index:])
            return _normalize_execution_plan(raw_plan, prompt)
        except json.JSONDecodeError:
            continue

    return _fallback_execution_plan(prompt)


async def build_prompt_with_history(prompt: str, mode: str, session_id: str) -> str:
    session = await session_store.get_session(mode, session_id)
    if session.thread.message_store is None:
        history_block = ""
    else:
        history_messages = await session.thread.message_store.list_messages()
        history_block = _format_history_messages(history_messages[-PROMPT_HISTORY_MESSAGES:])

    document_block = _format_uploaded_documents(session.uploaded_documents)
    if not history_block and not document_block:
        return prompt

    sections = []
    if history_block:
        sections.append(
            "Use the relevant information from the prior conversation when responding. "
            "If the history is not relevant, prioritize the latest request.\n\n"
            f"Conversation history:\n{history_block}"
        )
    if document_block:
        sections.append(
            "Use the uploaded documents when they are relevant to the latest request. "
            "If they are not relevant, do not force them into the answer.\n\n"
            f"{document_block}"
        )

    sections.append(f"Latest user request:\n{prompt}")
    return "\n\n".join(sections)


async def build_prompt_with_documents(prompt: str, mode: str, session_id: str) -> str:
    session = await session_store.get_session(mode, session_id)
    document_block = _format_uploaded_documents(session.uploaded_documents)
    if not document_block:
        return prompt

    return (
        "Use the uploaded documents when they are relevant to the latest request. "
        "If they are not relevant, do not force them into the answer.\n\n"
        f"{document_block}\n\n"
        f"Latest user request:\n{prompt}"
    )


async def append_session_exchange(mode: str, session_id: str, user_prompt: str, assistant_reply: str) -> None:
    if not assistant_reply.strip():
        return

    session = await session_store.get_session(mode, session_id)
    await session.thread.on_new_messages(
        [
            ChatMessage(role="user", text=user_prompt),
            ChatMessage(role="assistant", text=assistant_reply),
        ]
    )
    await session_store.save_session(mode, session_id)


async def generate_execution_plan(prompt: str, mode: str, session_id: str, model_chat_client) -> dict[str, object]:
    planner_agent = ChatAgent(
        name="Planner",
        chat_client=model_chat_client,
        instructions=get_text(
            'agent_planner_instructions',
            LANGUAGE,
            max_steps=PLAN_STEP_LIMIT,
            max_tools=PLAN_TOOL_LIMIT,
        ),
    )
    planner_prompt = await build_prompt_with_history(prompt, mode, session_id)
    planner_parts: list[str] = []
    async for update in planner_agent.run_stream(planner_prompt):
        if update.text:
            planner_parts.append(update.text)

    return _extract_plan_payload("".join(planner_parts), prompt)


def search_tool(
    query: Annotated[str, Field(description="Search query")],
) -> str:
    """Tool to search for guidelines in the financial sector"""
    started_at = time.time()
    try:
        if not query or not query.strip():
            _append_search_trace(
                {
                    "tool": "search_tool",
                    "query": "",
                    "status": "error",
                    "result_count": 0,
                    "duration_ms": round((time.time() - started_at) * 1000, 2),
                    "results": [],
                    "message": get_text('search_empty_query', LANGUAGE),
                }
            )
            return get_text('search_empty_query', LANGUAGE)
        
        # Initialize Azure AI Search client
        search_client = SearchClient(
            endpoint=SEARCH_ENDPOINT,
            index_name=SEARCH_INDEX_NAME,
            credential=AzureKeyCredential(SEARCH_API_KEY)
        )
        
        # Execute semantic search
        results = search_client.search(
            search_text=query,
            query_type=QueryType.SEMANTIC,
            semantic_configuration_name=SEARCH_SEMANTIC_CONFIG,
            top=3,
            select=["content", "metadata_storage_name"]
        )
        
        # Format results
        formatted_results = []
        structured_results = []
        file_label = get_text('search_file_label', LANGUAGE)
        content_label = get_text('search_content_label', LANGUAGE)
        
        for result in results:
            content = result.get("content", "")
            file_name = result.get("metadata_storage_name", "Unknown")
            if content and content.strip():  # Only add non-empty content
                formatted_results.append(f"{file_label}: {file_name}\n{content_label}: {content}\n")
                structured_results.append(
                    {
                        "source": _compact_text(file_name, limit=160),
                        "excerpt": _compact_text(content, limit=SEARCH_EXCERPT_LIMIT),
                    }
                )
        
        if formatted_results:
            logger.info(get_text('log_search_success', LANGUAGE, count=len(formatted_results)))
            _append_search_trace(
                {
                    "tool": "search_tool",
                    "query": _compact_text(query, limit=200),
                    "status": "success",
                    "result_count": len(structured_results),
                    "duration_ms": round((time.time() - started_at) * 1000, 2),
                    "results": structured_results,
                }
            )
            return "\n---\n".join(formatted_results)
        else:
            _append_search_trace(
                {
                    "tool": "search_tool",
                    "query": _compact_text(query, limit=200),
                    "status": "no_results",
                    "result_count": 0,
                    "duration_ms": round((time.time() - started_at) * 1000, 2),
                    "results": [],
                }
            )
            return get_text('search_no_results', LANGUAGE, query=query)
            
    except Exception as e:
        error_msg = get_text('search_error', LANGUAGE, error=str(e))
        logger.error(error_msg)
        _append_search_trace(
            {
                "tool": "search_tool",
                "query": _compact_text(query, limit=200),
                "status": "error",
                "result_count": 0,
                "duration_ms": round((time.time() - started_at) * 1000, 2),
                "results": [],
                "message": error_msg,
            }
        )
        return error_msg


@app.post("/api/stream")
async def api_stream(request: Request):
    start_time = time.time()
    request_id = f"req_{int(start_time * 1000)}"
    logger.info(f"[{request_id}] {get_text('log_request_received', LANGUAGE)}")
    
    body = await request.json()
    prompt = body.get("prompt", "")
    model_name = get_requested_model_name(body)
    session_id = get_session_id(body)

    logger.info(f"[{request_id}] {get_text('log_model_info', LANGUAGE, model=model_name)}")
    
    if not prompt:
        return JSONResponse({"error": "prompt required"}, status_code=400)
    
    parse_time = time.time()
    logger.info(f"[{request_id}] {get_text('log_request_parsed', LANGUAGE, time=f'{(parse_time - start_time)*1000:.2f}')}")

    async def generator():
        # Get a chat client for the requested model
        model_chat_client = get_chat_client_for_model(model_name)
        if model_chat_client is None:
            yield json.dumps({"type": "error", "message": get_text('error_config_missing', LANGUAGE)}, ensure_ascii=False) + "\n"
            return

        session = await session_store.get_session(MODE_GENERAL, session_id)
        async with session.lock:
            try:
                plan = await generate_execution_plan(prompt, MODE_GENERAL, session_id, model_chat_client)
                yield json.dumps({"type": "plan", "plan": plan}, ensure_ascii=False) + "\n"
                prompt_with_documents = await build_prompt_with_documents(prompt, MODE_GENERAL, session_id)

                # Create a simple agent
                agent_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_agent_creating', LANGUAGE)}")
                simple_agent = ChatAgent(
                    name="SimpleAgent",
                    chat_client=model_chat_client,
                    instructions=get_text('agent_simple_instructions', LANGUAGE),
                )
                logger.info(f"[{request_id}] {get_text('log_agent_created', LANGUAGE, time=f'{(time.time() - agent_start)*1000:.2f}')}")

                # Stream response
                stream_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_streaming_start', LANGUAGE, length=len(prompt))}")
                first_chunk = True
                chunk_count = 0
                async for update in simple_agent.run_stream(prompt_with_documents, thread=session.thread):
                    if update.text:
                        if first_chunk:
                            logger.info(f"[{request_id}] {get_text('log_first_chunk', LANGUAGE, time=f'{(time.time() - stream_start)*1000:.2f}')}")
                            first_chunk = False
                        chunk_count += 1
                        yield json.dumps({"type": "delta", "content": update.text}, ensure_ascii=False) + "\n"

                total_time = time.time() - start_time
                logger.info(f"[{request_id}] {get_text('log_completed', LANGUAGE, time=f'{total_time:.2f}', count=chunk_count)}")
                await session_store.save_session(MODE_GENERAL, session_id)
                yield json.dumps({"type": "complete"}, ensure_ascii=False) + "\n"
            except Exception as e:
                logger.exception("[%s] Regular chat stream failed", request_id)
                yield json.dumps({"type": "error", "message": str(e)}, ensure_ascii=False) + "\n"

    return StreamingResponse(generator(), media_type="application/x-ndjson")


@app.post("/api/rag/stream")
@app.post("/api/guideline/stream")
async def api_guideline_stream(request: Request):
    """Streaming endpoint for RAG-style referenced responses"""
    start_time = time.time()
    request_id = f"rag_{int(start_time * 1000)}"
    logger.info(f"[{request_id}] {get_text('log_guideline_request', LANGUAGE)}")
    
    body = await request.json()
    prompt = body.get("prompt", "")
    model_name = get_requested_model_name(body)
    session_id = get_session_id(body, "guideline")

    logger.info(f"[{request_id}] {get_text('log_model_info', LANGUAGE, model=model_name)}")
    
    if not prompt:
        return JSONResponse({"error": "prompt required"}, status_code=400)
    
    parse_time = time.time()
    logger.info(f"[{request_id}] {get_text('log_request_parsed', LANGUAGE, time=f'{(parse_time - start_time)*1000:.2f}')}")

    async def generator():
        # Get a chat client for the requested model
        model_chat_client = get_chat_client_for_model(model_name)
        if model_chat_client is None:
            yield json.dumps({"type": "error", "message": get_text('error_config_missing', LANGUAGE)}, ensure_ascii=False) + "\n"
            return

        session = await session_store.get_session(MODE_GUIDELINE, session_id)
        async with session.lock:
            trace_token = SEARCH_TRACE_CONTEXT.set([])
            try:
                plan = await generate_execution_plan(prompt, MODE_GUIDELINE, session_id, model_chat_client)
                yield json.dumps({"type": "plan", "plan": plan}, ensure_ascii=False) + "\n"
                prompt_with_documents = await build_prompt_with_documents(prompt, MODE_GUIDELINE, session_id)

                # Agent using Azure AI Search tool
                agent_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_search_agent_creating', LANGUAGE)}")
                search_agent = ChatAgent(
                    chat_client=model_chat_client,
                    instructions=get_text('agent_guideline_instructions', LANGUAGE),
                    tools=[search_tool],
                )
                logger.info(f"[{request_id}] {get_text('log_search_agent_created', LANGUAGE, time=f'{(time.time() - agent_start)*1000:.2f}')}")

                # Stream response
                stream_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_streaming_start', LANGUAGE, length=len(prompt))}")
                first_chunk = True
                chunk_count = 0
                async for update in search_agent.run_stream(prompt_with_documents, thread=session.thread):
                    if update.text:
                        if first_chunk:
                            logger.info(f"[{request_id}] {get_text('log_first_chunk', LANGUAGE, time=f'{(time.time() - stream_start)*1000:.2f}')}")
                            first_chunk = False
                        chunk_count += 1
                        yield json.dumps({"type": "delta", "content": update.text}, ensure_ascii=False) + "\n"

                search_traces = SEARCH_TRACE_CONTEXT.get() or []
                for trace in search_traces:
                    yield json.dumps({"type": "trace", "trace": trace}, ensure_ascii=False) + "\n"

                evidence_items = _collect_evidence_from_traces(search_traces)
                if evidence_items:
                    yield json.dumps({"type": "evidence", "evidence": evidence_items}, ensure_ascii=False) + "\n"

                total_time = time.time() - start_time
                logger.info(f"[{request_id}] {get_text('log_completed', LANGUAGE, time=f'{total_time:.2f}', count=chunk_count)}")
                await session_store.save_session(MODE_GUIDELINE, session_id)
                yield json.dumps({"type": "complete"}, ensure_ascii=False) + "\n"

            except Exception as e:
                error_msg = get_text('error_search_processing', LANGUAGE, error=str(e))
                logger.error(f"[{request_id}] ❌ {error_msg}")
                yield json.dumps({"type": "error", "message": error_msg}, ensure_ascii=False) + "\n"
            finally:
                SEARCH_TRACE_CONTEXT.reset(trace_token)

    return StreamingResponse(generator(), media_type="application/x-ndjson")


@app.post("/api/multi-agent-stream")
async def multi_agent_stream(request: Request):
    start_time = time.time()
    request_id = f"multi_{int(start_time * 1000)}"
    logger.info(f"[{request_id}] {get_text('log_multi_agent_request', LANGUAGE)}")
    
    body = await request.json()
    prompt = body.get("prompt", "")
    model_name = get_requested_model_name(body)
    session_id = get_session_id(body)

    logger.info(f"[{request_id}] {get_text('log_model_info', LANGUAGE, model=model_name)}")
    
    if not prompt:
        return JSONResponse({"error": "prompt required"}, status_code=400)
    
    parse_time = time.time()
    logger.info(f"[{request_id}] {get_text('log_multi_agent_parsed', LANGUAGE, time=f'{(parse_time - start_time)*1000:.2f}')}")

    async def generator():
        # Get a chat client for the requested model
        model_chat_client = get_chat_client_for_model(model_name)
        if model_chat_client is None:
            yield json.dumps(
                {"type": "error", "message": get_text('error_config_missing', LANGUAGE)},
                ensure_ascii=False,
            ) + "\n"
            return
        
        session = await session_store.get_session(MODE_GENERAL, session_id)
        async with session.lock:
            # Create model-specific agents
            model_critical_agent = ChatAgent(
                name="CriticalAnalyst",
                chat_client=model_chat_client,
                instructions=get_text('agent_critical_instructions', LANGUAGE),
            )

            model_positive_agent = ChatAgent(
                name="PositiveAdvocate",
                chat_client=model_chat_client,
                instructions=get_text('agent_positive_instructions', LANGUAGE),
            )

            model_synthesizer_agent = ChatAgent(
                name="Synthesizer",
                chat_client=model_chat_client,
                instructions=get_text('agent_synthesizer_instructions', LANGUAGE),
            )

            try:
                yield json.dumps(
                    {"type": "ui_message", "message": get_text('ui_multi_agent_start', LANGUAGE)},
                    ensure_ascii=False,
                ) + "\n"
                yield json.dumps({"type": "start"}, ensure_ascii=False) + "\n"

                workflow_prompt = await build_prompt_with_history(prompt, MODE_GENERAL, session_id)
                plan = await generate_execution_plan(prompt, MODE_GENERAL, session_id, model_chat_client)
                yield json.dumps({"type": "plan", "plan": plan}, ensure_ascii=False) + "\n"

                # Use ConcurrentBuilder to create parallel workflow
                workflow_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_workflow_building', LANGUAGE)}")
                agents = [model_critical_agent, model_positive_agent]
                workflow = ConcurrentBuilder().participants(agents).build()
                logger.info(f"[{request_id}] {get_text('log_workflow_built', LANGUAGE, time=f'{(time.time() - workflow_start)*1000:.2f}')}")

                # Store agent results
                agent_results = {}

                # Execute workflow with streaming
                workflow_exec_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_parallel_execution_start', LANGUAGE, length=len(prompt))}")
                event_count = 0
                async for event in workflow.run_stream(workflow_prompt):
                    event_count += 1
                    # AgentRunUpdateEvent: Streaming updates from agents
                    if isinstance(event, AgentRunUpdateEvent):
                        agent_name = event.executor_id

                        if event.data and hasattr(event.data, 'text') and event.data.text:
                            data = {
                                "agent": agent_name,
                                "content": event.data.text,
                                "is_final": False
                            }
                            yield json.dumps(data, ensure_ascii=False) + "\n"

                    elif isinstance(event, WorkflowOutputEvent):
                        if event.data:
                            for msg in event.data:
                                if hasattr(msg, 'author_name') and hasattr(msg, 'text'):
                                    agent_results[msg.author_name] = msg.text

                agents_time = time.time() - workflow_exec_start
                logger.info(f"[{request_id}] {get_text('log_parallel_execution_complete', LANGUAGE, time=f'{agents_time:.2f}', count=event_count)}")
                yield json.dumps({"type": "agents_complete"}, ensure_ascii=False) + "\n"

                # Stream synthesis analysis
                synthesis_start = time.time()
                logger.info(f"[{request_id}] {get_text('log_synthesis_start', LANGUAGE)}")
                yield json.dumps({"type": "synthesis_start"}, ensure_ascii=False) + "\n"

                critical_content = agent_results.get("CriticalAnalyst", "")
                positive_content = agent_results.get("PositiveAdvocate", "")

                synthesis_prompt = f"""
Integrate the following two perspectives to provide a balanced analysis.

Original question: {prompt}

Critical perspective:
{critical_content}

Positive perspective:
{positive_content}
"""

                synthesis_parts = []
                synthesis_chunk_count = 0
                async for update in model_synthesizer_agent.run_stream(synthesis_prompt):
                    if update.text:
                        synthesis_chunk_count += 1
                        synthesis_parts.append(update.text)
                        data = {
                            "agent": "Synthesizer",
                            "content": update.text,
                            "is_final": False
                        }
                        yield json.dumps(data, ensure_ascii=False) + "\n"

                synthesis_text = "".join(synthesis_parts).strip()
                assistant_summary = (
                    "Critical perspective:\n"
                    f"{critical_content}\n\n"
                    "Positive perspective:\n"
                    f"{positive_content}\n\n"
                    "Synthesis:\n"
                    f"{synthesis_text}"
                )
                await append_session_exchange(MODE_GENERAL, session_id, prompt, assistant_summary)

                synthesis_time = time.time() - synthesis_start
                total_time = time.time() - start_time
                logger.info(f"[{request_id}] {get_text('log_synthesis_complete', LANGUAGE, time=f'{synthesis_time:.2f}', count=synthesis_chunk_count)}")
                logger.info(f"[{request_id}] {get_text('log_overall_complete', LANGUAGE, time=f'{total_time:.2f}')}")
                yield json.dumps({"type": "complete"}, ensure_ascii=False) + "\n"

            except Exception as e:
                import traceback
                error_data = {
                    "type": "error",
                    "message": str(e),
                    "traceback": traceback.format_exc()
                }
                yield json.dumps(error_data, ensure_ascii=False) + "\n"

    return StreamingResponse(generator(), media_type="application/x-ndjson")


@app.post("/api/phase1/stream")
async def phase1_planning_stream(request: Request):
    start_time = time.time()
    request_id = f"phase1_{int(start_time * 1000)}"
    logger.info(f"[{request_id}] {get_text('log_idobata_request', LANGUAGE)}")

    body = await request.json()
    prompt = body.get("prompt", "")
    model_name = get_requested_model_name(body)
    tone = body.get("tone", "balanced")
    session_id = get_session_id(body, "idobata")

    logger.info(f"[{request_id}] {get_text('log_model_info', LANGUAGE, model=model_name)}")
    logger.info(f"[{request_id}] {get_text('log_tone_setting', LANGUAGE, tone=tone)}")

    if not prompt:
        return JSONResponse({"error": "prompt required"}, status_code=400)

    parse_time = time.time()
    logger.info(f"[{request_id}] {get_text('log_request_parsed', LANGUAGE, time=f'{(parse_time - start_time)*1000:.2f}')}")
    
    # Tone-specific speaking style instructions
    tone_instructions = {
        "formal": get_text('tone_formal', LANGUAGE),
        "balanced": get_text('tone_balanced', LANGUAGE),
        "casual": get_text('tone_casual', LANGUAGE),
        "concise": get_text('tone_concise', LANGUAGE),
        "detailed": get_text('tone_detailed', LANGUAGE)
    }
    
    tone_suffix = tone_instructions.get(tone, tone_instructions["balanced"])

    async def generator():
        model_chat_client = get_chat_client_for_model(model_name)
        if model_chat_client is None:
            yield json.dumps({"type": "error", "message": get_text('error_config_missing', LANGUAGE)}, ensure_ascii=False) + "\n"
            return

        AGENT_SEQUENCE = [
            "CEO",
            "CTO",
            "CFO",
            "COO",
        ]

        def _pick_participant_name(*candidates: Optional[str]) -> Optional[str]:
            """Return the first candidate that matches a known participant name."""
            for candidate in candidates:
                if not candidate:
                    continue
                value = str(candidate).strip()
                if value in AGENT_SEQUENCE:
                    return value
            return None

        session = await session_store.get_session(MODE_IDOBATA, session_id)
        async with session.lock:
            # AI Board Meeting: Create CxO agents
            logger.info(f"[{request_id}] {get_text('log_planning_agent_creating', LANGUAGE)}")
            ceo_agent = ChatAgent(
                name="CEO",
                chat_client=model_chat_client,
                description="CEO leading the management meeting and presenting strategic direction",
                instructions=get_text('agent_board_ceo_instructions', LANGUAGE, tone_suffix=tone_suffix),
            )

            cto_agent = ChatAgent(
                name="CTO",
                chat_client=model_chat_client,
                description="Evaluates technical strategy and feasibility",
                instructions=get_text('agent_board_cto_instructions', LANGUAGE, tone_suffix=tone_suffix),
            )

            cfo_agent = ChatAgent(
                name="CFO",
                chat_client=model_chat_client,
                description="Evaluates financial viability and business potential",
                instructions=get_text('agent_board_cfo_instructions', LANGUAGE, tone_suffix=tone_suffix),
            )

            coo_agent = ChatAgent(
                name="COO",
                chat_client=model_chat_client,
                description="Integrates CxO opinions and creates execution plan",
                instructions=get_text('agent_board_coo_instructions', LANGUAGE, tone_suffix=tone_suffix),
            )

            logger.info(f"[{request_id}] {get_text('log_planning_agent_created', LANGUAGE)}")

            # Build workflow (create new instance to reset)
            workflow_id = f"wf_{int(time.time() * 1000)}"
            logger.info(f"[{request_id}] {get_text('log_board_workflow_building', LANGUAGE, id=workflow_id)}")

            selector_state = {
                "calls": 0,
                "last_returned": None,
            }

            def planning_selector(state) -> Optional[str]:
                round_idx = getattr(state, "current_round", None)
                if round_idx is None:
                    round_idx = getattr(state, "round_index", 0) or 0

                history = getattr(state, "conversation", None)
                if history is None:
                    history = getattr(state, "history", ()) or ()

                MAX_ROUNDS = 10
                if round_idx >= MAX_ROUNDS:
                    logger.warning(f"[{request_id}] {get_text('warning_max_rounds', LANGUAGE, max=MAX_ROUNDS)}")
                    return None

                MAX_SELECTOR_CALLS = 50
                if selector_state["calls"] >= MAX_SELECTOR_CALLS:
                    logger.warning(f"[{request_id}] {get_text('warning_max_selector_calls', LANGUAGE, max=MAX_SELECTOR_CALLS)}")
                    return None

                for turn in reversed(history):
                    text = ""
                    if hasattr(turn, "text"):
                        text = turn.text or ""
                    elif hasattr(turn, "message") and hasattr(turn.message, "text"):
                        text = turn.message.text or ""
                    if text and "PLAN_READY:" in text:
                        logger.info(f"[{request_id}] {get_text('log_plan_ready', LANGUAGE)}")
                        return None

                last_agent_speaker: Optional[str] = None
                for turn in reversed(history):
                    speaker = None
                    if hasattr(turn, "author_name"):
                        speaker = turn.author_name
                    elif hasattr(turn, "speaker"):
                        speaker = turn.speaker
                    elif hasattr(turn, "message") and hasattr(turn.message, "author_name"):
                        speaker = turn.message.author_name
                    speaker = _pick_participant_name(speaker)
                    if speaker:
                        last_agent_speaker = speaker
                        break

                if not last_agent_speaker:
                    next_speaker = AGENT_SEQUENCE[0]
                else:
                    current_index = AGENT_SEQUENCE.index(last_agent_speaker)
                    next_speaker = AGENT_SEQUENCE[(current_index + 1) % len(AGENT_SEQUENCE)]

                if next_speaker == selector_state.get("last_returned"):
                    current_index = AGENT_SEQUENCE.index(next_speaker)
                    next_speaker = AGENT_SEQUENCE[(current_index + 1) % len(AGENT_SEQUENCE)]

                selector_state["calls"] += 1
                selector_state["last_returned"] = next_speaker
                return next_speaker

            planning_workflow = (
                GroupChatBuilder()
                .participants([
                    ceo_agent,
                    cto_agent,
                    cfo_agent,
                    coo_agent,
                ])
                .with_select_speaker_func(planning_selector, orchestrator_name="PlanningOrchestrator")
                .build()
            )

            logger.info(f"[{request_id}] {get_text('log_board_workflow_built', LANGUAGE, id=workflow_id)}")

            workflow_prompt = await build_prompt_with_history(prompt, MODE_IDOBATA, session_id)
            plan = await generate_execution_plan(prompt, MODE_IDOBATA, session_id, model_chat_client)
            yield json.dumps({"type": "plan", "plan": plan}, ensure_ascii=False) + "\n"
            workflow_exec_start = time.time()
            logger.info(f"[{request_id}] {get_text('log_board_workflow_start', LANGUAGE, length=len(prompt))}")
            event_count = 0
            seq = 0
            board_notes = {
                "CEO": [],
                "CTO": [],
                "CFO": [],
                "COO": [],
            }

            async for event in planning_workflow.run_stream(workflow_prompt):
                event_count += 1
                if isinstance(event, AgentRunUpdateEvent):
                    author_name = None
                    if event.data is not None:
                        if hasattr(event.data, "author_name") and event.data.author_name:
                            author_name = event.data.author_name
                        elif hasattr(event.data, "message") and hasattr(event.data.message, "author_name"):
                            author_name = event.data.message.author_name

                    agent_name = _pick_participant_name(author_name, event.executor_id)
                    if not agent_name:
                        continue

                    text = ""
                    if event.data is not None and hasattr(event.data, "text"):
                        text = event.data.text or ""
                    elif event.data is not None and hasattr(event.data, "message") and hasattr(event.data.message, "text"):
                        text = event.data.message.text or ""
                    elif event.data is not None:
                        text = str(event.data)
                    if text:
                        board_notes[agent_name].append(text)
                        seq += 1
                        data = {
                            "agent": agent_name,
                            "seq": seq,
                            "content": text,
                            "is_final": False,
                            "executor_id": getattr(event, "executor_id", None),
                            "author_name": author_name,
                        }
                        yield json.dumps(data, ensure_ascii=False) + "\n"
                elif isinstance(event, WorkflowOutputEvent):
                    pass

            assistant_summary = (
                "CEO:\n"
                f"{''.join(board_notes['CEO']).strip()}\n\n"
                "CTO:\n"
                f"{''.join(board_notes['CTO']).strip()}\n\n"
                "CFO:\n"
                f"{''.join(board_notes['CFO']).strip()}\n\n"
                "COO:\n"
                f"{''.join(board_notes['COO']).strip()}"
            )
            await append_session_exchange(MODE_IDOBATA, session_id, prompt, assistant_summary)

            total_time = time.time() - start_time
            workflow_time = time.time() - workflow_exec_start
            logger.info(f"[{request_id}] {get_text('log_board_complete', LANGUAGE, workflow_time=f'{workflow_time:.2f}', count=event_count, total_time=f'{total_time:.2f}')}")
            yield json.dumps({"type": "complete"}, ensure_ascii=False) + "\n"

    return StreamingResponse(generator(), media_type="application/x-ndjson")


@app.get("/api/models")
async def get_models():
    return {
        "default_model": DEFAULT_MODEL_NAME,
        "models": get_model_metadata(),
    }


@app.get("/api/files")
async def list_uploaded_files(session_id: str = "default", mode: str = MODE_GENERAL):
    session = await session_store.get_session(mode, session_id)
    async with session.lock:
        return {"files": [_document_metadata(document) for document in session.uploaded_documents]}


@app.post("/api/files")
async def upload_files(
    files: list[UploadFile] = File(...),
    session_id: str = Form("default"),
    mode: str = Form(MODE_GENERAL),
):
    if not files:
        return JSONResponse({"error": get_text('upload_error_missing_files', LANGUAGE)}, status_code=400)

    session = await session_store.get_session(mode, session_id)
    async with session.lock:
        if len(session.uploaded_documents) + len(files) > MAX_UPLOAD_FILES:
            return JSONResponse(
                {"error": get_text('upload_error_too_many_files', LANGUAGE, limit=MAX_UPLOAD_FILES)},
                status_code=400,
            )

        uploaded: list[dict[str, Any]] = []
        for upload in files:
            filename = upload.filename or "uploaded-file"
            extension = _guess_upload_extension(filename)
            if extension not in ALLOWED_UPLOAD_EXTENSIONS:
                return JSONResponse(
                    {"error": get_text('upload_error_unsupported_type', LANGUAGE, filename=filename)},
                    status_code=400,
                )

            data = await upload.read()
            if len(data) > MAX_UPLOAD_FILE_SIZE_BYTES:
                return JSONResponse(
                    {
                        "error": get_text(
                            'upload_error_file_too_large',
                            LANGUAGE,
                            filename=filename,
                            limit_mb=f"{MAX_UPLOAD_FILE_SIZE_BYTES / (1024 * 1024):.1f}",
                        )
                    },
                    status_code=400,
                )

            try:
                extracted_text = _extract_document_text(filename, data)
            except Exception as exc:
                return JSONResponse({"error": str(exc)}, status_code=400)

            normalized_text = extracted_text.strip()
            if not normalized_text:
                return JSONResponse(
                    {"error": get_text('upload_error_no_text', LANGUAGE, filename=filename)},
                    status_code=400,
                )

            normalized_text = normalized_text[:MAX_UPLOAD_TEXT_CHARS]
            document = UploadedDocument(
                file_id=uuid4().hex,
                name=filename,
                content_type=upload.content_type or "application/octet-stream",
                size_bytes=len(data),
                extracted_text=normalized_text,
                preview_text=_compact_text(normalized_text, limit=220),
            )
            session.uploaded_documents.append(document)
            uploaded.append(_document_metadata(document))

        await session_store.save_session(mode, session_id)
        return {
            "files": uploaded,
            "all_files": [_document_metadata(document) for document in session.uploaded_documents],
        }


@app.delete("/api/files/{file_id}")
async def delete_uploaded_file(file_id: str, session_id: str = "default", mode: str = MODE_GENERAL):
    session = await session_store.get_session(mode, session_id)
    async with session.lock:
        before_count = len(session.uploaded_documents)
        session.uploaded_documents = [
            document for document in session.uploaded_documents if document.file_id != file_id
        ]
        if len(session.uploaded_documents) == before_count:
            return JSONResponse({"error": get_text('upload_error_file_not_found', LANGUAGE)}, status_code=404)

        await session_store.save_session(mode, session_id)
        return {"status": "ok", "files": [_document_metadata(document) for document in session.uploaded_documents]}


@app.post("/api/sessions/clear")
async def clear_session(request: Request):
    body = await request.json()
    session_id = get_session_id(body)
    mode = str(body.get("mode", MODE_GENERAL)).strip() or MODE_GENERAL
    await session_store.clear_session(mode, session_id)
    return {"status": "ok"}


@app.get("/")
async def root():
    return {"status": "ok", "framework": "Microsoft Agent Framework"}

if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", "8000"))
    host = os.getenv("HOST", "0.0.0.0")
    print("=== Agent Framework (Workflow Edition) Backend Starting ===")
    print(f"URL: http://localhost:{port}")
    print(f"Docs: http://localhost:{port}/docs")
    print("\nCtrl+C to stop\n")
    uvicorn.run(app, host=host, port=port)
